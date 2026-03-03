import os
import re
from docx import Document
from openpyxl import load_workbook
from config import G4Config,LABEL_RE,REQ_ID_RE


# ============================================================
# BASIC READERS
# ============================================================

def read_all_text(docx_path):
    doc = Document(docx_path)
    blocks = []

    for p in doc.paragraphs:
        if p.text.strip():
            blocks.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    blocks.append(cell.text.strip())

    return blocks


def read_docx(path):
    return Document(path)


def extract_docx_text(path):
    doc = Document(path)
    parts = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)

    return "\n".join(parts)


# ============================================================
# G4 EXTRACTORS
# ============================================================

def extract_from_docx(file_path):
    doc = Document(file_path)
    requirements = []

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) >= 2:
                rid = row.cells[0].text.strip()
                rtxt = row.cells[1].text.strip()

                if rid.startswith(G4Config.REQ_PREFIX):
                    requirements.append((rid, rtxt))

    return requirements


def extract_from_excel(file_path):
    wb = load_workbook(file_path, data_only=True)
    requirements = []

    for sheet in wb.worksheets:
        for row in sheet.iter_rows(min_col=1, max_col=2):
            rid, rtxt = row[0].value, row[1].value
            if isinstance(rid, str) and rid.startswith(G4Config.REQ_PREFIX):
                requirements.append((rid.strip(), str(rtxt).strip()))

    return requirements


def collect_requirements(input_dir):
    """
    Discovers files and extracts requirements.
    NO review logic here.
    """
    all_requirements = []

    for fname in os.listdir(input_dir):
        if fname.startswith("~$"):
            continue
        if not ("SRS" in fname.upper() or "TEST" in fname.upper()):
            continue

        path = os.path.join(input_dir, fname)

        if fname.lower().endswith(".docx"):
            all_requirements.extend(extract_from_docx(path))
        elif fname.lower().endswith(".xlsx"):
            all_requirements.extend(extract_from_excel(path))

    return all_requirements


# ============================================================
# NORMALIZATION
# ============================================================

def clean(text):
    return text.replace("\n", " ").strip() if text else ""


def norm(text):
    """
    Soft normalization for similarity functions.
    Does NOT delete symbols like >= in the raw pipeline.
    """
    return re.sub(r"\s+", " ", str(text).lower().strip()) if text else ""


def normalize_text(text):
    """
    Strong NLP normalization.
    WARNING: removes punctuation, so use ONLY for similarity/intent matching.
    """
    if not text:
        return ""
    text = re.sub(r"[^\w\s]", " ", str(text).lower())
    return re.sub(r"\s+", " ", text).strip()


# ============================================================
# REQUIREMENT EXTRACTION (G5/G7 compatibility)
# ============================================================

VALID_ID_PATTERN = re.compile(r"^\s*(SCU_STC_SRS_\d+)\b", re.IGNORECASE)


def _extract_requirements_from_text(text):
    """
    Extract requirements from flattened text.
    Returns:
        dict { REQ_ID : normalized requirement text }
    Used by G1.4
    """
    from config import REQ_ID_RE

    reqs = {}
    lines = [l.strip() for l in str(text).splitlines() if l.strip()]
    i = 0

    while i < len(lines):
        m = REQ_ID_RE.match(lines[i])
        if m:
            rid = m.group(1)
            lookahead = " ".join(lines[i:i + 5]).lower()

            if re.search(r"\b(shall|should|must)\b", lookahead):
                buf = []
                i += 1
                while i < len(lines) and not REQ_ID_RE.match(lines[i]):
                    buf.append(lines[i])
                    i += 1
                reqs[rid] = normalize_text(" ".join(buf))
                continue
        i += 1

    return reqs


def _extract_requirements_from_docx(docx_path):
    """
    Extract requirements from DOCX tables.
    Returns:
        list[dict]: [{"id":..., "raw_id":..., "text":...}]
    """
    doc = Document(docx_path)
    requirements = []

    for table in doc.tables:
        for row in table.rows:
            if len(row.cells) < 2:
                continue

            raw_id = row.cells[0].text.strip()
            text = row.cells[1].text.strip()

            if raw_id.lower() in ("requirement id", "srs id", "id"):
                continue

            if not raw_id and text:
                requirements.append({"id": None, "raw_id": "", "text": text})
                continue

            if raw_id and not VALID_ID_PATTERN.match(raw_id):
                requirements.append({"id": None, "raw_id": raw_id, "text": text})
                continue

            match = VALID_ID_PATTERN.match(raw_id)
            requirements.append({
                "id": match.group(1),
                "raw_id": raw_id,
                "text": text
            })

    return requirements


def extract_requirements(arg):
    """
    Single public function name supports BOTH call styles.
    """
    # Heuristic: if it's a .docx file path, treat as DOCX
    if isinstance(arg, str) and arg.strip().lower().endswith(".docx") and os.path.exists(arg):
        return _extract_requirements_from_docx(arg)

    # Otherwise treat as raw text
    return _extract_requirements_from_text(arg)


# ============================================================
# G1 WRAPPER (RAW + NORM safe)
# ============================================================

SYS_ID_PATTERN = re.compile(r"[A-Z0-9_]+_SYS_\d+", re.IGNORECASE)
HLR_ID_PATTERN = re.compile(r"[A-Z0-9_]+_SRS_\d+", re.IGNORECASE)


class G1IOUtils:
    """Safe IO wrapper for G1."""

    @staticmethod
    def read_docx_tables(path):
        doc = Document(path)
        rows = []
        for table in doc.tables:
            for row in table.rows:
                rows.append([cell.text.strip() for cell in row.cells])
        return rows

    @staticmethod
    def is_arinc_trigger(text):
        return "arinc" in str(text).lower() and "label" in str(text).lower()

    @staticmethod
    def is_arinc_table(table_rows):
        if not table_rows or not table_rows[0]:
            return False
        return "label" in table_rows[0][0].lower()

    @staticmethod
    def read_docx_tables_grouped(path):
        doc = Document(path)
        all_tables = []

        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                table_rows.append([cell.text.strip() for cell in row.cells])
            all_tables.append(table_rows)

        return all_tables

    @staticmethod
    def extract_arinc_from_cell(cell):
        labels = []

        for tbl in cell.tables:
            header = tbl.rows[0].cells[0].text.lower()
            if "label" not in header:
                continue

            for row in tbl.rows[1:]:
                val = row.cells[0].text.strip()
                if val:
                    labels.append(val)

        return labels

    @staticmethod
    def extract_cell_text_and_tables(cell):
        """
        Preserve newlines in RAW to support state-diagram + IF/ELSE parsing.
        """
        text_parts = []
        table_blocks = []

        for p in cell.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text.strip())

        for table in cell.tables:
            rows = []
            for row in table.rows:
                row_text = " | ".join(
                    c.text.strip() for c in row.cells if c.text.strip()
                )
                if row_text:
                    rows.append(row_text)
            if rows:
                table_blocks.append("\n".join(rows))

        return "\n".join(text_parts), "\n\n".join(table_blocks)

    @staticmethod
    def build_system_text_to_id_map(system_reqs):
        mapping = {}
        for r in system_reqs:
            mapping[r.get("TEXT_RAW", r.get("TEXT", ""))] = r["SYS_ID"]
        return mapping

    @staticmethod
    def normalize_id(x: str) -> str:
        return re.sub(r"[^A-Z0-9_]", "", (x or "").upper())

    @staticmethod
    def load_system_requirements(path):
        doc = Document(path)
        system_reqs = []

        for table in doc.tables:
            for row in table.rows[1:]:
                if len(row.cells) < 2:
                    continue

                sys_cell = row.cells[0].text.replace("\u00a0", " ").strip()
                desc_cell = row.cells[1]

                if sys_cell.lower() in ("id", "sys id", "requirement id"):
                    continue

                m = SYS_ID_PATTERN.search(sys_cell)
                if not m:
                    continue
                sys_id = G1IOUtils.normalize_id(m.group(0))

                raw_text, table_text = G1IOUtils.extract_cell_text_and_tables(desc_cell)

                system_reqs.append({
                    "SYS_ID": sys_id,
                    "TEXT_RAW": raw_text,
                    "TEXT": raw_text,  # backward compatible
                    "TEXT_NORM": normalize_text(raw_text),
                    "TABLE_TEXT": table_text
                })

        return system_reqs

    @staticmethod
    def load_software_requirements(path):
        doc = Document(path)
        hlr_reqs = []

        for table in doc.tables:
            for row in table.rows[1:]:
                if len(row.cells) < 2:
                    continue

                hlr_cell = row.cells[0].text.replace("\u00a0", " ").strip()
                desc_cell = row.cells[1]

                m = HLR_ID_PATTERN.search(hlr_cell)
                if not m:
                    continue
                hlr_id = m.group(0)

                raw_text, table_text = G1IOUtils.extract_cell_text_and_tables(desc_cell)

                hlr_reqs.append({
                    "HLR_ID": hlr_id,
                    "TEXT_RAW": raw_text,
                    "TEXT": raw_text,  # backward compatible
                    "TEXT_NORM": normalize_text(raw_text),
                    "TABLE_TEXT": table_text
                })

        return hlr_reqs

    @staticmethod
    def load_traceability_srs_to_sys(path):
        rows = G1IOUtils.read_docx_tables(path)
        trace_links = []

        for row in rows[1:]:
            if len(row) < 3:
                continue

            hlr_id = G1IOUtils.normalize_id(row[0])
            sys_id = G1IOUtils.normalize_id(row[2])


            if not hlr_id:
                continue

            trace_links.append({
                "HLR_ID": hlr_id,
                "SYS_ID": sys_id if sys_id else "NOT_TRACED"
            })

        return trace_links

    @staticmethod
    def normalize_text(text):
        # keep your class-specific normalizer if needed
        if not text:
            return ""
        return " ".join(str(text).split())
    
class G1_3_4IOUtils:
    # io_utils.py
    # Combined utilities for G1.3 + G1.4 (FINAL)

    # =========================================================
    # BASIC TEXT HELPERS (USED BY G1.3)
    # =========================================================
    def clean(text):
        return text.replace("\n", " ").strip()

    def norm(text):
        return re.sub(r"\s+", " ", str(text).lower().strip())

    def extract_label(text):
        m = LABEL_RE.search(text or "")
        return m.group(1) if m else None

    def parse_bits(text):
        nums = list(map(int, re.findall(r"\d+", str(text))))
        if not nums:
            return None, None
        if len(nums) == 1:
            return nums[0], nums[0]
        return min(nums), max(nums)

    # =========================================================
    # G1.4 – FULL DOCX TEXT EXTRACTION
    # =========================================================
    def extract_docx_text(path):
        """
        Extracts ALL text from a DOCX:
        - paragraphs
        - table cells
        """
        doc = Document(path)
        parts = []

        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())

        return "\n".join(parts)

    # =========================================================
    # G1.4 – REQUIREMENT EXTRACTION (HLR / SYS)
    # =========================================================
    def extract_requirements(text):
        """
        Extracts requirements in format:
        SCU_xxx_ID → normalized text
        """
        reqs = {}
        lines = [l.strip() for l in text.splitlines() if l.strip()]
        i = 0

        while i < len(lines):
            m = REQ_ID_RE.match(lines[i])
            if m:
                rid = m.group(1)
                lookahead = " ".join(lines[i:i+5]).lower()

                if re.search(r"\b(shall|should|must)\b", lookahead):
                    buf = []
                    i += 1
                    while i < len(lines) and not REQ_ID_RE.match(lines[i]):
                        buf.append(lines[i])
                        i += 1
                    reqs[rid] = norm(" ".join(buf))
                    continue
            i += 1

        return reqs

    # =========================================================
    # G1.4 – LABEL EXTRACTION FROM TEXT
    # =========================================================
    def extract_labels(text):
        return {
            m for m in LABEL_RE.findall(text or "")
            if m.isdigit() and 0 < int(m) <= 377
        }
