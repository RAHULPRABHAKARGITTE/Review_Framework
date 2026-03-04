# main_g4.py

import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

import io_utils
import g4.g4_logic


# ================= PATHS =================
def run_g4():

    # Get project root dynamically (Review_Framework folder)
    PROJECT_ROOT = Path(__file__).resolve().parent.parent

    INPUT_DIR = PROJECT_ROOT / "inputs"
    OUTPUT_DIR = PROJECT_ROOT / "outputs"
    OUTPUT_FILE = OUTPUT_DIR / "CI_G4_output.xlsx"

    # Create outputs folder if not exists
    OUTPUT_DIR.mkdir(exist_ok=True)

    # ================= READ INPUTS =================
    all_requirements = io_utils.collect_requirements(str(INPUT_DIR))

    # ================= EXCEL OUTPUT =================
    wb = Workbook()
    ws = wb.active
    ws.title = "HLR's are verifiable"

    # ---- Title ----
    ws.merge_cells("A1:K1")
    ws["A1"] = "High-Level Requirements are Verifiable"
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    # ---- Headers ----
    headers = [
        "Req ID",
        "G 4.1,G4.4:\nCheck testability for each requirement and Ensure verification is practical within project constraints.",
        "G 4.2:\nDefine acceptance criteria: Ensure measurable and objective criteria are provided.",
        "G 4.2 Failure Reason",
        "G 4.3 Verification Method (Declared)",
        "G 4.5 Use mandatory language",
        "Register Requirement",
        "Communication Requirement",
        "Fault Requirement",
        "I/O Requirement",
        "Functional Requirement"
    ]

    ws.append(headers)

    # ---- Header Formatting ----
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A3"

    # ================= DATA =================
    for rid, rtxt in all_requirements:

        g42_result, g42_reason = g4.g4_logic.check_g42_acceptance_criteria(rtxt)

        ws.append([
            rid,
            g4.g4_logic.check_g41_testability(rtxt),
            g42_result,
            g42_reason,
            g4.g4_logic.extract_verification_method(rtxt),
            g4.g4_logic.check_g45(rtxt),
            "Yes" if g4.g4_logic.is_register_requirement(rtxt) else "No",
            "Yes" if g4.g4_logic.is_communication_requirement(rtxt) else "No",
            "Yes" if g4.g4_logic.is_fault_requirement(rtxt) else "No",
            "Yes" if g4.g4_logic.is_io_requirement(rtxt) else "No",
            "Yes" if g4.g4_logic.is_functional_requirement(rtxt) else "No",
        ])

    # ================= ALIGNMENT =================
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row):
        for cell in row:
            if cell.column == 1:
                # Req ID LEFT aligned
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ================= AUTO COLUMN WIDTH =================
    for col in ws.columns:
        max_len = 0
        column_letter = get_column_letter(col[0].column)

        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(max_len + 2, 60)

    # ================= SAVE FILE =================
    wb.save(str(OUTPUT_FILE))

    print("\n✅ G4 review completed successfully")
    print("📄 Output:", OUTPUT_FILE)


if __name__ == "__main__":
    run_g4()