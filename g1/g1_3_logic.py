# g1_3_logic.py
import os
import re
from docx import Document
import pandas as pd

from config import CommonConfig, G1Config
from io_utils import G1_3_4IOUtils

# =========================================================
# STEP 1: Generate Software_req.docx from SRS
# =========================================================

SRS_ID_RE = re.compile(r"SCU_STC_SRS_\d+")


def _in_path(fname: str) -> str:
    return os.path.join(CommonConfig.BASE_INPUT, fname)


def _out_path(fname: str) -> str:
    return os.path.join(CommonConfig.BASE_OUTPUT, fname)


def generate_software_req():
    """
    Generates a normalized Software_req.docx by extracting requirement rows from the SRS.
    Output location: outputs/<G1Config.SOFTWARE_REQ_DOCX>
    """
    os.makedirs(CommonConfig.BASE_OUTPUT, exist_ok=True)

    out_path = _out_path(G1Config.SOFTWARE_REQ_DOCX)

    if os.path.exists(out_path):
        try:
            os.remove(out_path)
        except PermissionError:
            raise RuntimeError(f"❌ Please close '{out_path}' and rerun.")

    src = Document(_in_path(G1Config.SRS_DOCX))
    out = Document()

    table = out.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "ID"
    table.rows[0].cells[1].text = "DESC"

    for tbl in src.tables:
        for row in tbl.rows:
            text = " ".join(c.text for c in row.cells)
            m = SRS_ID_RE.search(text)
            if not m:
                continue

            rid = m.group()
            r = table.add_row()
            r.cells[0].text = rid
            dst = r.cells[1]

            for cell in row.cells:
                if rid in cell.text:
                    continue

                for p in cell.paragraphs:
                    if p.text.strip():
                        dst.add_paragraph(p.text.strip())

                for nt in cell.tables:
                    t = dst.add_table(rows=len(nt.rows), cols=len(nt.columns))
                    t.style = "Table Grid"
                    for i, rr in enumerate(nt.rows):
                        for j, cc in enumerate(rr.cells):
                            t.rows[i].cells[j].text = cc.text.strip()

    out.save(out_path)
    print(f"✅ Generated {out_path}")
    return out_path


def extract_srs_arinc(software_req_docx: str):
    doc = Document(software_req_docx)
    rx, msg = [], []

    for tbl in doc.tables:
        for row in tbl.rows[1:]:
            req_id = row.cells[0].text.strip()
            cell = row.cells[1]
            label = G1_3_4IOUtils.extract_label(cell.text)

            for nt in cell.tables:
                headers = [G1_3_4IOUtils.norm(c.text) for c in nt.rows[0].cells]

                # RX RATE
                if "receiver" in " ".join(headers):
                    for r in nt.rows[1:]:
                        c = [G1_3_4IOUtils.clean(x.text) for x in r.cells]
                        if len(c) >= 3:
                            rx.append({
                                "Req ID": req_id,
                                "Label": G1_3_4IOUtils.extract_label(c[1]) or label,
                                "Receiver": c[0],
                                "Interval": c[2],
                            })

                # MESSAGE DEFINITION
                if "bit" in " ".join(headers):
                    for r in nt.rows[1:]:
                        c = [G1_3_4IOUtils.clean(x.text) for x in r.cells]
                        if len(c) >= 4:
                            msg.append({
                                "Req ID": req_id,
                                "Label": label,
                                "Bit": c[0],
                                "Field": c[1],
                                "Description": c[2],
                                "Definition": c[3],
                            })

    return rx, msg


def extract_icd_arinc():
    doc = Document(_in_path(G1Config.ICD_DOCX))

    rx, msg = [], []
    current_label = None

    tables = list(doc.tables)
    ti = 0

    for p in doc.paragraphs:
        lbl = G1_3_4IOUtils.extract_label(p.text)
        if lbl:
            current_label = lbl

        if ti < len(tables) and tables[ti]._tbl.getprevious() == p._p:
            t = tables[ti]
            ti += 1

            headers = [G1_3_4IOUtils.norm(c.text) for c in t.rows[0].cells]

            if "receiver" in " ".join(headers):
                for r in t.rows[1:]:
                    c = [G1_3_4IOUtils.clean(x.text) for x in r.cells]
                    if len(c) >= 3:
                        rx.append({
                            "Label": G1_3_4IOUtils.extract_label(c[1]),
                            "Receiver": c[0],
                            "Interval": c[2],
                        })

            if "bit" in " ".join(headers):
                for r in t.rows[1:]:
                    c = [G1_3_4IOUtils.clean(x.text) for x in r.cells]
                    if len(c) >= 4:
                        msg.append({
                            "Label": current_label,
                            "Bit": c[0],
                            "Field": c[1],
                            "Description": c[2],
                            "Definition": c[3],
                        })

    return rx, msg


def compare():
    """
    Main entrypoint for G1.3.
    Returns: (msg_out, rx_out)
    """
    software_req_docx = generate_software_req()

    srs_rx, srs_msg = extract_srs_arinc(software_req_docx)
    icd_rx, icd_msg = extract_icd_arinc()

    msg_out, rx_out = [], []

    for s in srs_msg:
        match = next(
            (i for i in icd_msg
             if i.get("Label") == s.get("Label")
             and G1_3_4IOUtils.norm(i.get("Field")) == G1_3_4IOUtils.norm(s.get("Field"))),
            None
        )

        if not match:
            msg_out.append({**s, "Status": "Fail", "Reason": "Field not found in ICD"})
            continue

        sb1, _ = G1_3_4IOUtils.parse_bits(s.get("Bit"))
        ib1, _ = G1_3_4IOUtils.parse_bits(match.get("Bit"))

        fails = []
        if sb1 != (ib1 + 1 if ib1 is not None else None):
            fails.append("Bit offset mismatch (SRS = ICD + 1)")
        if G1_3_4IOUtils.norm(s.get("Description")) != G1_3_4IOUtils.norm(match.get("Description")):
            fails.append("Description mismatch")
        if G1_3_4IOUtils.norm(s.get("Definition")) != G1_3_4IOUtils.norm(match.get("Definition")):
            fails.append("Definition mismatch")

        msg_out.append({
            **s,
            "ICD Bit": match.get("Bit", ""),
            "Status": "Pass" if not fails else "Fail",
            "Reason": " | ".join(fails)
        })

    for s in srs_rx:
        match = next(
            (i for i in icd_rx
             if i.get("Label") == s.get("Label")
             and G1_3_4IOUtils.norm(i.get("Receiver")) == G1_3_4IOUtils.norm(s.get("Receiver"))),
            None
        )

        if not match:
            rx_out.append({**s, "Status": "Fail", "Reason": "Receiver not found in ICD"})
            continue

        ok = G1_3_4IOUtils.norm(s.get("Interval")) == G1_3_4IOUtils.norm(match.get("Interval"))
        rx_out.append({
            **s,
            "ICD Interval": match.get("Interval", ""),
            "Status": "Pass" if ok else "Fail",
            "Reason": "" if ok else "Transmission interval mismatch"
        })

    return msg_out, rx_out
