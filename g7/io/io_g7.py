# io_utils.py
import os
import pandas as pd
from docx import Document
from openpyxl import load_workbook


def read_srs_doc(path: str):
    """
    Read the DOCX SRS and return a list of table rows (each row is a list of cell texts).
    """
    doc = Document(path)
    rows = []
    for t in doc.tables:
        for r in t.rows:
            rows.append([c.text.strip() for c in r.cells])
    return rows


def write_requirements_excel(
        df: pd.DataFrame,
        algo_df: pd.DataFrame,
        non_algo_df: pd.DataFrame,
        outfile: str,
        all_req_sheet: str,
        algo_req_sheet: str,
        non_algo_req_sheet: str
):
    """
    Write three dataframes to an Excel workbook in separate sheets.
    Overwrites the file using xlsxwriter.
    """
    with pd.ExcelWriter(outfile, engine="xlsxwriter") as writer:
        df.to_excel(writer, sheet_name=all_req_sheet, index=False)
        algo_df.to_excel(writer, sheet_name=algo_req_sheet, index=False)
        non_algo_df.to_excel(writer, sheet_name=non_algo_req_sheet, index=False)


def append_results(
        excel_file: str,
        summary_df: pd.DataFrame,
        details_df: pd.DataFrame,
        summary_sheet: str,
        details_sheet: str
):
    """
    Append (or create) results summary and details dataframes to a workbook.
    If file doesn't exist → create with headers.
    If exists → append to end of each sheet without headers.
    """
    if not os.path.exists(excel_file):
        with pd.ExcelWriter(excel_file, mode="w", engine="xlsxwriter") as writer:
            summary_df.to_excel(writer, sheet_name=summary_sheet, index=False, header=True)
            details_df.to_excel(writer, sheet_name=details_sheet, index=False, header=True)
        print(f"Created {excel_file} with sheets '{summary_sheet}' and '{details_sheet}'")
        return

    book = load_workbook(excel_file)

    def _append_df(writer, sheet_name, df):
        if sheet_name in book.sheetnames:
            ws = book[sheet_name]
            startrow = ws.max_row  # append below last row
            df.to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=startrow)
        else:
            df.to_excel(writer, sheet_name=sheet_name, index=False, header=True)

    with pd.ExcelWriter(excel_file, mode="a", engine="openpyxl") as writer:
        _append_df(writer, summary_sheet, summary_df)
        _append_df(writer, details_sheet, details_df)

    print(f"Results appended to {excel_file} → '{summary_sheet}' and '{details_sheet}'")
