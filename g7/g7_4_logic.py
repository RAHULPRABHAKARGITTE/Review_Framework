
import os
import re
from typing import List, Dict, Tuple, Optional
import pandas as pd
from docx import Document

# -----------------------------
# Configuration
# -----------------------------
INPUT_DOCX_PATH = r"C:/Ali/SCU_SRS.docx"
OUTPUT_XLSX_PATH = "SCU_SRS_Performance_Review.xlsx"

# Objective kept for output; Project ID extracted but NOT written anywhere
TARGET_OBJECTIVE_TEXT = "Objective G 7.4"
TARGET_ID_TEXT = "SCU_STC_SRS"

CHECKPOINTS = [
    "Timing - WCET Defined",
    "Timing - Deadlines / Response Time",
    "Timing - Task Rates",
    "Resource - CPU Utilization",
    "Resource - Memory Constraints",
    "Resource - I/O Bandwidth",
]

# Heuristic regex patterns to extract quantifications
TIME_UNITS = r"(ns|us|µs|microseconds?|ms|milliseconds?|s|sec|seconds?)"
FREQ_UNITS = r"(Hz|kHz|MHz|cycles/s|updates?/s)"
SIZE_UNITS = r"(B|KB|KiB|MB|MiB|GB|GiB|bytes?)"
BANDWIDTH_UNITS = r"(bps|kbps|Mbps|Gbps|B/s|KB/s|MB/s|GB/s)"
PERCENT = r"(\d+(\.\d+)?)\s*%"

# -----------------------------
# Utilities
# -----------------------------

def normalize_text(s: str) -> str:
    return re.sub(r"\s+", " ", s or "").strip()

def excerpt_with_context(text: str, start: int, end: int, ctx: int = 80) -> str:
    left = max(0, start - ctx)
    right = min(len(text), end + ctx)
    snippet = text[left:right]
    return normalize_text(snippet)

def find_with_units(text: str, patterns: List[Tuple[str, str]]) -> Tuple[str, Optional[str]]:
    """
    Search text using list of (label, regex_pattern).
    Returns first matching (evidence_snippet, extracted_value).
    """
    for _, pat in patterns:
        for m in re.finditer(pat, text, flags=re.IGNORECASE):
            span = m.span()
            snippet = excerpt_with_context(text, span[0], span[1])
            value = m.group(0)
            return snippet, value
    return "", None

def status_from_match(value: Optional[str], quantified_patterns: List[str], raw_text: str) -> str:
    if not value:
        return "Missing"
    # If a number with plausible units appears, treat as quantified
    for qp in quantified_patterns:
        if re.search(qp, value, flags=re.IGNORECASE):
            return "Present (Quantified)"
    # If units not in captured value but present in context, still count as quantified
    for qp in quantified_patterns:
        if re.search(qp, raw_text, flags=re.IGNORECASE):
            return "Present (Quantified)"
    return "Present (Unquantified)"

# -----------------------------
# Metadata Extraction (Objective kept; Project ID not output)
# -----------------------------

def get_full_text_from_doc(doc: Document) -> str:
    parts = []
    for p in doc.paragraphs:
        parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return normalize_text("\n".join(parts))

def extract_metadata_from_doc(doc: Document) -> Tuple[str, str]:
    """
    Extract:
      - Objective Id: prefer exact 'Objective G 7.4'; fallback to 'Objective <G 7.4>' pattern; default TARGET_OBJECTIVE_TEXT
      - Project ID: prefer 'SCU_STC_SRS' (accepting space/hyphen/underscore variants); or from 'ID: <value>'; default TARGET_ID_TEXT
    Note: Project ID is extracted but will NOT be written to the output per request.
    """
    text = get_full_text_from_doc(doc)

    # Objective extraction
    objective = None
    if re.search(r"\bObjective\s*G\s*7\.4\b", text, flags=re.IGNORECASE):
        objective = TARGET_OBJECTIVE_TEXT
    else:
        m = re.search(r"\bObjective(?:\s*(?:Id|ID|Identifier))?\s*[:\-]?\s*([A-Za-z]\s*\d+(?:\.\d+)+)\b", text, flags=re.IGNORECASE)
        if m:
            val = normalize_text(m.group(1))
            objective = f"Objective {val}"
    if not objective:
        objective = TARGET_OBJECTIVE_TEXT

    # Project ID extraction (not output)
    project_id = None
    if re.search(r"\bSCU[\s_\-]*STC[\s_\-]*SRS\b", text, flags=re.IGNORECASE):
        project_id = TARGET_ID_TEXT
    else:
        m2 = re.search(r"\bID\b\s*[:\-]\s*([A-Za-z0-9._\-\s]+)", text, flags=re.IGNORECASE)
        if m2:
            cand = normalize_text(m2.group(1)).upper()
            cand = re.sub(r"[\s\-]+", "_", cand)
            project_id = cand
    if not project_id:
        project_id = TARGET_ID_TEXT

    return objective, project_id

# -----------------------------
# Requirement Extraction (First Column Only)
# -----------------------------

HEADER_HINTS = ["id", "req", "requirement", "identifier"]

def is_header_row(cells_text: List[str]) -> bool:
    header_join = " | ".join([c.lower() for c in cells_text])
    return any(hint in header_join for hint in HEADER_HINTS)

def extract_requirements_from_first_col(doc: Document) -> List[Dict]:
    """
    Extract requirement IDs exclusively from the FIRST COLUMN of every table.
    The rest of the row (columns 2..N) is concatenated for analysis (not shown in output).
    """
    reqs = []
    for t_index, t in enumerate(doc.tables):
        rows = []
        for row in t.rows:
            rows.append([normalize_text(cell.text) for cell in row.cells])

        if not rows:
            continue

        for r_index, cells in enumerate(rows):
            if not cells:
                continue

            # Skip header-like first row
            if r_index == 0 and is_header_row(cells):
                continue

            req_id = normalize_text(cells[0]) if len(cells) >= 1 else ""
            if not req_id:
                continue  # first column must carry the Requirement ID

            # Skip if first column is a header label
            if req_id.lower() in ("id", "req id", "requirement id", "identifier"):
                continue

            # Concatenate the rest of the row as analysis text (not exported as a column)
            other_text = normalize_text(" | ".join(c for c in cells[1:])) if len(cells) > 1 else ""

            reqs.append({
                "id": req_id,
                "row_text": other_text,     # used internally for checkpoint detection
                "table_index": t_index,
                "row_index": r_index,
                "source": "table-first-col"
            })

    # De-duplicate by Requirement ID, retaining the first occurrence
    seen = set()
    unique = []
    for r in reqs:
        if r["id"] not in seen:
            unique.append(r)
            seen.add(r["id"])
    return unique

def extract_all_requirements(doc_path: str) -> Tuple[List[Dict], Tuple[str, str]]:
    doc = Document(doc_path)
    objective_id, project_id = extract_metadata_from_doc(doc)
    reqs = extract_requirements_from_first_col(doc)
    return reqs, (objective_id, project_id)

# -----------------------------
# Checkpoint Detectors (use row_text)
# -----------------------------

def detect_wcet(text: str) -> Tuple[str, Optional[str], str]:
    patterns = [
        ("wcet", r"\b(WCET|worst[-\s]?case\s+execution\s+time)\b.*?\b\d+(\.\d+)?\s*"+TIME_UNITS),
        ("wcet_phrase", r"\b(WCET|worst[-\s]?case\s+execution\s+time)\b"),
    ]
    snippet, value = find_with_units(text, patterns)
    status = status_from_match(value, [TIME_UNITS], text)
    return snippet, value, status

def detect_deadline_response(text: str) -> Tuple[str, Optional[str], str]:
    patterns = [
        ("deadline", r"\b(deadline|latency|response\s*time)\b.*?\b\d+(\.\d+)?\s*"+TIME_UNITS),
        ("deadline_phrase", r"\b(deadline|latency|response\s*time)\b"),
    ]
    snippet, value = find_with_units(text, patterns)
    status = status_from_match(value, [TIME_UNITS], text)
    return snippet, value, status

def detect_task_rates(text: str) -> Tuple[str, Optional[str], str]:
    patterns = [
        ("frequency", r"\b(\d+(\.\d+)?)\s*"+FREQ_UNITS),
        ("period", r"\b(period|cycle|update\s*rate)\b.*?\b\d+(\.\d+)?\s*"+TIME_UNITS),
        ("rate_phrase", r"\b(rate|period|frequency|update\s*rate|cycle)\b"),
    ]
    snippet, value = find_with_units(text, patterns)
    status = status_from_match(value, [TIME_UNITS, FREQ_UNITS], text)
    return snippet, value, status

def detect_cpu_util(text: str) -> Tuple[str, Optional[str], str]:
    patterns = [
        ("cpu_percent", r"\b(CPU|processor)\b.*?"+PERCENT),
        ("utilization", r"\b(utilization|usage|load)\b.*?"+PERCENT),
        ("phrase", r"\b(CPU\s*utilization|CPU\s*usage|processor\s*load)\b"),
    ]
    snippet, value = find_with_units(text, patterns)
    status = status_from_match(value, [PERCENT], text)
    return snippet, value, status

def detect_memory(text: str) -> Tuple[str, Optional[str], str]:
    patterns = [
        ("memory_size", r"\b(memory|RAM|ROM|flash|stack|heap)\b.*?\b\d+(\.\d+)?\s*"+SIZE_UNITS),
        ("stack_heap", r"\b(stack|heap)\b.*?\b\d+(\.\d+)?\s*"+SIZE_UNITS),
        ("phrase", r"\b(memory|RAM|ROM|flash|stack|heap)\b"),
    ]
    snippet, value = find_with_units(text, patterns)
    status = status_from_match(value, [SIZE_UNITS], text)
    return snippet, value, status

def detect_io_bandwidth(text: str) -> Tuple[str, Optional[str], str]:
    patterns = [
        ("bw_numeric", r"\b(bandwidth|throughput)\b.*?\b\d+(\.\d+)?\s*"+BANDWIDTH_UNITS),
        ("net_units", r"\b\d+(\.\d+)?\s*"+BANDWIDTH_UNITS+r"\b"),
        ("phrase", r"\b(bandwidth|throughput|I/O|bus)\b"),
    ]
    snippet, value = find_with_units(text, patterns)
    status = status_from_match(value, [BANDWIDTH_UNITS], text)
    return snippet, value, status

DETECTOR_MAP = {
    "Timing - WCET Defined": detect_wcet,
    "Timing - Deadlines / Response Time": detect_deadline_response,
    "Timing - Task Rates": detect_task_rates,
    "Resource - CPU Utilization": detect_cpu_util,
    "Resource - Memory Constraints": detect_memory,
    "Resource - I/O Bandwidth": detect_io_bandwidth,
}

# -----------------------------
# Review + Excel Export
# -----------------------------

def review_requirements(reqs: List[Dict], meta: Tuple[str, str]) -> pd.DataFrame:
    objective_id, _project_id = meta  # project_id extracted but not used in output
    rows = []
    for r in reqs:
        r_id = r["id"]
        row_text = r.get("row_text", "")
        for chk in CHECKPOINTS:
            detector = DETECTOR_MAP[chk]
            snippet, value, status = detector(row_text)
            comment = ""
            if status == "Missing":
                if "Timing" in chk:
                    comment = "No explicit timing value found; specify quantified targets with units."
                elif "CPU" in chk:
                    comment = "No CPU utilization target; specify a max % or headroom."
                elif "Memory" in chk:
                    comment = "No memory constraint; specify RAM/ROM/stack/heap footprints."
                elif "I/O" in chk:
                    comment = "No bandwidth/throughput limit; define kbps/MB/s with margins."
            elif status == "Present (Unquantified)":
                comment = "Constraint mentioned but lacks numeric target/units; add measurable criteria."

            rows.append({
                "Objective_ID": objective_id,
                "Requirement_ID": r_id,
                "Checkpoint": chk,
                "Status": status,
                "Evidence_Snippet": snippet,
                "Extracted_Value": value if value else "",
                "Comments": comment
            })
    df = pd.DataFrame(rows)
    return df

def write_excel(df: pd.DataFrame, out_path: str):
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        # Only the Detailed Review sheet (no Metadata, no Summary)
        df.to_excel(writer, index=False, sheet_name="Detailed Review")

        # Column widths – Detailed Review
        ws = writer.sheets["Detailed Review"]
        col_widths = {
            "A": 18,  # Objective_ID
            "B": 28,  # Requirement_ID
            "C": 34,  # Checkpoint
            "D": 22,  # Status
            "E": 70,  # Evidence_Snippet
            "F": 24,  # Extracted_Value
            "G": 60,  # Comments
        }
        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

    print(f"✅ Excel written to: {os.path.abspath(out_path)}")

def main():
    if not os.path.exists(INPUT_DOCX_PATH):
        print(f"❌ Input file not found: {INPUT_DOCX_PATH}")
        print("Please verify the path or update INPUT_DOCX_PATH.")
        return

    print("📄 Loading SRS document...")
    reqs, meta = extract_all_requirements(INPUT_DOCX_PATH)
    objective_id, project_id = meta
    ##print(f"🔖 Extracted Objective Id: '{objective_id}' (Project ID found: '{project_id}', not written to output)")

    if not reqs:
        print("⚠️ No requirement IDs found in the SRS Document. Check the document format.")
        return

    print(f"🔎 Extracted {len(reqs)} unique Requirement IDs from SRS Document. Reviewing against checkpoints...")
    df = review_requirements(reqs, meta)
    write_excel(df, OUTPUT_XLSX_PATH)

if __name__ == "__main__":
    main()
