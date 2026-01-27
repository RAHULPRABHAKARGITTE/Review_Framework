# -*- coding: utf-8 -*-
"""
DO-178C-Oriented SRS Requirement Reviewer (Table-first, ID in Column 1)
-----------------------------------------------------------------------

Reads a Word SRS (.docx), extracts requirements from tables where:
  - Column 1 = Requirement ID
  - Next non-empty column in the row = Requirement Text

Assesses each requirement against:
  - Correctness
  - Determinism
  - Testability
  - Robustness (ALWAYS PASS per current user requirement)
  - Verifiability

Special handling per user request:
  - Atomicity ("Requirement may not be atomic...") does NOT mark Correctness as False.

Outputs an Excel file WITHOUT:
  - Requirement_Text column
  - Source/Source_Index columns (not created at all)
  - Extra sheets (Checkpoint Summary / Overall / Reviewer Guidance)

Author: (You)
"""

import os
import re
from collections import defaultdict
from typing import List, Dict, Optional

import pandas as pd
from docx import Document

# ------------------------------- Configuration -------------------------------

INPUT_DOCX_PATH = r"C:/Ali/SCU_SRS.docx"
OUTPUT_XLSX_PATH = None  # if None, will be derived from input, e.g., ..._DO178C_Review.xlsx

# If your requirement text is not in the second column, set a fallback search range here
# Example: [1, 2] means try column 1 first (0-based index), then 2
TEXT_COLUMN_PREFERENCE = None  # None => "first non-empty cell after column 0"

# Do NOT count atomicity as an error for Correctness
ATOMICITY_COUNTS_AS_ERROR = False
# If True, we’ll still add a suggestion line (but not an error) to help authors
ADD_ATOMICITY_SUGGESTION = True

# Words/phrases to flag as ambiguous or non-deterministic
AMBIGUOUS_TERMS = [
    r'\betc\.?\b', r'\band\/or\b', r'\bas appropriate\b', r'\bas needed\b',
    r'\bif possible\b', r'\bas soon as possible\b', r'\buser-friendly\b',
    r'\bintuitive\b', r'\boptimize\b', r'\bapproximately\b', r'\babout\b',
    r'\broughly\b', r'\btypical(ly)?\b', r'\busually\b', r'\bgenerally\b',
    r'\bfast\b', r'\bquick(ly)?\b', r'\bsoon\b', r'\bshould\b', r'\bmay\b'
]

# Common TBD/TBR placeholders
PLACEHOLDER_TERMS = [r'\bTBD\b', r'\bTBR\b', r'\bTBC\b', r'\bTBS\b', r'XXX', r'\?\?\?']

# Phrases suggesting acceptance criteria or verification method
ACCEPTANCE_PHRASES = [
    r'\bwithin\b', r'\bno more than\b', r'\bat least\b', r'\bnot more than\b',
    r'\bless than\b', r'\bgreater than\b', r'\bequal to\b', r'\b±', r'\bplus\/minus\b',
    r'\bshall be verified\b', r'\bverified by\b', r'\bverification\b',
    r'\btest(ed)?\b', r'\banalysis\b', r'\breview\b', r'\binspection\b'
]

# Robustness / error / boundary terms
ROBUSTNESS_TERMS = [
    r'\berror\b', r'\bfault\b', r'\bfail(-| )?safe\b', r'\bfault(-| )?tolerant\b',
    r'\bboundary\b', r'\blimit(s)?\b', r'\bmin(imum)?\b', r'\bmax(imum)?\b', r'\brange\b',
    r'\binvalid input(s)?\b', r'\bout of range\b', r'\bexception\b', r'\brecovery\b',
    r'\bdegraded\b', r'\bgraceful\b', r'\bdefault behavior\b', r'\bsaturation\b', r'\boverflow\b'
]

# Units or tokens that often indicate measurability
UNITS_OR_NUMERIC = [
    r'\d+\s*(ms|s|sec|seconds|Hz|kHz|MHz|%|ppm|°C|deg|degrees|mA|A|V|mV|g|kg|N|bit|bits|byte|bytes|KB|MB|GB|px|m|cm|mm|km|in|ft|fps|kPa|Pa|bar)\b',
    r'\d+(\.\d+)?\b'  # any number
]

# Sentence split regex for basic heuristics
SENTENCE_SPLIT = re.compile(r'(?<=[.!?])\s+')


# -----------------------------------------------------------------------------


def load_docx(path: str) -> Document:
    if not os.path.exists(path):
        raise FileNotFoundError(f"Input file not found: {path}")
    return Document(path)


def _match_any(patterns: List[str], text: str, flags=re.IGNORECASE) -> bool:
    return any(re.search(p, text, flags=flags) for p in patterns)


def normalize(text: str) -> str:
    return re.sub(r'\s+', ' ', text.strip())


def extract_requirements_from_tables_col1_ids(doc: Document) -> List[Dict]:
    """
    Extract requirements by assuming:
      - First column = Requirement ID (non-empty & not a header)
      - Requirement Text = first non-empty cell after column 0
        (or by TEXT_COLUMN_PREFERENCE if provided).
    Deduplicates by (ID, Text).
    """
    results = []
    seen = set()

    for tbl in doc.tables:
        for row in tbl.rows:
            cells = row.cells
            if len(cells) == 0:
                continue

            rid = normalize(cells[0].text) if cells[0].text else ""
            # Skip empty IDs and likely header rows
            if not rid or rid.lower() in {"id", "requirement id", "req id", "requirement"}:
                continue

            # Determine requirement text cell
            req_text = ""
            if TEXT_COLUMN_PREFERENCE:
                for ci in TEXT_COLUMN_PREFERENCE:
                    if ci < len(cells):
                        t = normalize(cells[ci].text)
                        if t:
                            req_text = t
                            break
            else:
                # First non-empty cell after column 0
                for ci in range(1, len(cells)):
                    t = normalize(cells[ci].text)
                    if t:
                        req_text = t
                        break

            if not req_text:
                continue

            key = (rid, req_text.lower())
            if key in seen:
                continue
            seen.add(key)

            results.append({
                "Requirement_ID": rid,
                "Requirement_Text": req_text,
            })

    return results


def _has_measurable_numbers(text: str) -> bool:
    return _match_any(UNITS_OR_NUMERIC, text)


def _has_acceptance_phrases(text: str) -> bool:
    return _match_any(ACCEPTANCE_PHRASES, text)


def _has_ambiguous_terms(text: str) -> bool:
    return _match_any(AMBIGUOUS_TERMS, text)


def _has_placeholders(text: str) -> bool:
    return _match_any(PLACEHOLDER_TERMS, text)


def analyze_requirement(req_text: str) -> Dict:
    """
    Assess the requirement text against the five checkpoints and generate
    a list of identified issues and suggestions.

    NOTE: Atomicity ("Requirement may not be atomic...") does not make Correctness False
    per user request. We only add an optional suggestion.

    NOTE: Per user request, Robustness always passes and never adds the
    "Robustness not addressed..." issue.
    """
    text = normalize(req_text)
    lower = text.lower()
    issues = []
    suggestions = []

    # --- Correctness ---
    correctness_pass = True
    if not re.search(r'\b(shall|must)\b', lower):
        correctness_pass = False
        issues.append("Missing prescriptive modal ('shall' or 'must').")
        suggestions.append("Rewrite using 'shall' for mandatory behavior (e.g., 'The system shall ...').")

    if _has_placeholders(lower):
        correctness_pass = False
        issues.append("Contains placeholder(s) (TBD/TBR/TBC/XXX/???).")
        suggestions.append("Resolve all placeholders with finalized values and references.")

    # Atomicity heuristic — DO NOT mark correctness as False (per request)
    sentences = SENTENCE_SPLIT.split(text) if text else []
    long_sentence = any(len(s) > 280 for s in sentences) if sentences else False
    too_many_conjunctions = any(
        len(re.findall(r'\band\b|\bor\b|,', s, flags=re.IGNORECASE)) >= 3 for s in sentences) if sentences else False
    if (long_sentence or too_many_conjunctions):
        if ATOMICITY_COUNTS_AS_ERROR:
            # (kept for toggling; default False)
            correctness_pass = False
            issues.append("Requirement may not be atomic (very long or compound).")
        if ADD_ATOMICITY_SUGGESTION:
            suggestions.append("Consider splitting into smaller, atomic requirements (one verifiable statement each).")

    # --- Determinism ---
    determinism_pass = True
    if _has_ambiguous_terms(lower):
        determinism_pass = False
        issues.append("Contains ambiguous or non-deterministic terms (e.g., 'etc.', 'and/or', 'approximately').")
        suggestions.append("Replace vague terms with precise, quantifiable statements and explicit conditions.")

    if re.search(r'\b(latency|response|deadline|period|rate|frequency|time)\b', lower) and not _has_measurable_numbers(
            lower):
        determinism_pass = False
        issues.append("Timing referenced without numeric bounds.")
        suggestions.append("Specify timing quantitatively (e.g., 'within 50 ms', 'period = 10 ms ±1 ms').")

    # --- Testability ---
    testability_pass = True
    measurable = _has_measurable_numbers(lower) or _has_acceptance_phrases(lower)
    if not measurable:
        testability_pass = False
        issues.append("Lacks measurable acceptance criteria.")
        suggestions.append("Add measurable thresholds, ranges, or explicit verification criteria.")

    # --- Robustness (ALWAYS PASS as per user request) ---
    robustness_pass = True
    # Intentionally do not add any robustness-related issues/suggestions.

    # --- Verifiability ---
    verifiability_pass = True
    if not measurable or _has_ambiguous_terms(lower):
        verifiability_pass = False
        if not measurable:
            issues.append("Verifiability risk: lacks quantifiable criteria.")
            suggestions.append("State quantitative criteria and verification method (Test/Analysis/Inspection/Review).")
        else:
            issues.append("Verifiability risk: contains ambiguous language.")
            suggestions.append("Eliminate ambiguous terms to enable unambiguous verification.")
    if not re.search(r'\b(test|analysis|inspection|review)\b', lower):
        suggestions.append("Add an explicit verification method tag, e.g., 'Verification: Test'.")

    passes = {
        "Correctness": correctness_pass,
        "Determinism": determinism_pass,
        "Testability": testability_pass,
        "Robustness": robustness_pass,
        "Verifiability": verifiability_pass
    }
    pass_count = sum(1 for v in passes.values() if v)
    fail_count = 5 - pass_count
    overall = pass_count == 5

    severity = "High" if fail_count >= 3 else ("Medium" if fail_count == 2 else ("Low" if fail_count == 1 else "None"))

    issues = sorted(set(issues))
    suggestions = sorted(set(suggestions))

    return {
        **passes,
        "Pass_Count": pass_count,
        "Fail_Count": fail_count,
        "Overall_Pass": overall,
        "Severity": severity,
        "Issues": "; ".join(issues) if issues else "",
        "Suggested_Fixes": "; ".join(suggestions) if suggestions else ""
    }


def review_srs(input_path: str, output_path: Optional[str] = None) -> str:
    doc = load_docx(input_path)
    requirements = extract_requirements_from_tables_col1_ids(doc)

    if not requirements:
        raise RuntimeError(
            "No requirements found. Ensure your SRS tables have Requirement IDs in the first column "
            "and requirement text in the second (or any subsequent) column."
        )

    rows = []
    for r in requirements:
        analysis = analyze_requirement(r["Requirement_Text"])
        row = {
            "Requirement_ID": r["Requirement_ID"],
            # DO NOT include Requirement_Text in output (per request)
            **analysis
        }
        rows.append(row)

    df = pd.DataFrame(rows)
    # Order columns WITHOUT Requirement_Text and WITHOUT Source/Source_Index
    col_order = [
        "Requirement_ID",
        "Correctness", "Determinism", "Testability", "Robustness", "Verifiability",
        "Pass_Count", "Fail_Count", "Overall_Pass", "Severity", "Issues", "Suggested_Fixes"
    ]
    df = df[col_order]

    if not output_path:
        base, _ = os.path.splitext(input_path)
        output_path = base + "_DO178C_Review.xlsx"

    # Write ONLY the main sheet using openpyxl engine
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Requirement Review", index=False)

    return output_path


if __name__ == "__main__":
    try:
        out_path = review_srs(INPUT_DOCX_PATH, OUTPUT_XLSX_PATH)
        print(f"✅ Review complete. Output saved to: {out_path}")
    except Exception as e:
        print(f"❌ Error: {e}")

