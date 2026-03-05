import pandas as pd
import os
from g2.g2_2_logic import check_g2_2_and_generate_excel
from g2.g2_3_logic import read_icd, check_g2_3_and_generate_excel
from g2.g2_5_logic import analyze_requirement_g2_5
from g2.g2_6_logic import check_g2_6_and_generate_excel
from io_utils import G2_IOUtils
from config import G2Config,CommonConfig
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.cell import MergedCell
import io_utils
import g2.g2_4_7_logic
from pathlib import Path


def combine_g2_excels(output_dir):
    """
    Combine all individual G2 Excel outputs into
    one single Excel file with multiple sheets.

    Output:
        G2_Result.xlsx
            - Sheet: G2_2
            - Sheet: G2_3
            - Sheet: G2_5
            - Sheet: G2_6
    """

    #g2_output_path = os.path.join(output_dir, "G2_Result.xlsx")

    # File paths (must match your existing filenames)
    file_map = {
        "G2_2": "G2_2_Consistency.xlsx",
        "G2_3": "G2_3_Terminology_and_Units.xlsx",
        "G2_5": "G2_5_Performance_Constraints.xlsx",
        "G2_6": "G2_6_Findings.xlsx"
    }

    with pd.ExcelWriter(str(G2Config.G2_OUTPUT), engine="openpyxl") as writer:

        for sheet_name, file_name in file_map.items():
            file_path = os.path.join(output_dir, file_name)

            if os.path.exists(file_path):
                df = pd.read_excel(file_path)
                df.to_excel(writer, sheet_name=sheet_name, index=False)
            else:
                # If file not found, create empty placeholder sheet
                empty_df = pd.DataFrame({
                    "Info": [f"{file_name} not found"]
                })
                empty_df.to_excel(writer, sheet_name=sheet_name, index=False)

    print(f"\n✅ Combined Excel created at:\n{str(G2Config.G2_OUTPUT)}")

    return str(G2Config.G2_OUTPUT)

def G2_2_logic():
    srs_path_g2_2 = str( G2Config.G2_SRS)
    G2_2_output_folder = str( G2Config.OUTPUT_DIR)

    # Step 1: Run G2.2
    excel_path_g2_2 = check_g2_2_and_generate_excel(srs_path_g2_2, G2_2_output_folder)

    print("G2.2 check completed")
    print("Excel generated at:", excel_path_g2_2)

def G2_3_logic():
    srs_path = str( G2Config.G2_SRS)
    icd_path = str( G2Config.ICD)
    G2_3_output_folder = str( G2Config.OUTPUT_DIR)

    # Step 1: Read ICD
    read_icd(icd_path)

    # Step 2: Read SRS (your unified reader)
    requirements = G2_IOUtils.read_srs_requirements(srs_path, mode="AUTO")

    # Step 3: Run G2.3
    df, excel_path = check_g2_3_and_generate_excel(requirements, G2_3_output_folder)

    print("G2.3 check completed")
    print("Excel generated at:", excel_path)

def G2_5_logic(srs_path, datasheet_path, output_folder):
    # -------------------------
    # Read SRS requirements
    # -------------------------
    requirements = G2_IOUtils.read_srs_requirements(srs_path)

    # -------------------------
    # Datasheet references (initial/static – can be enhanced later)
    # -------------------------
    # NOTE: You are passing datasheet_path for future PDF parsing;
    # for now we extract only key performance constraints
    datasheet_refs = {
        "cpu frequency": "150 mhz",
        "interrupt latency": "12 cycles",
        "flash wait state": "1",
        "adc conversion time": "80 ns"
    }

    # -------------------------
    # Run G2.5 analysis
    # -------------------------
    all_issues = []

    for req_id, text in requirements.items():
        issues = analyze_requirement_g2_5(req_id, text, datasheet_refs)
        all_issues.extend(issues)

    # -------------------------
    # Handle PASS case
    # -------------------------
    if not all_issues:
        all_issues.append({
            "Requirement_ID": "N/A",
            "Category": "PASS",
            "Violated_Content": "N/A",
            "Explanation": "No G2.5 violations found"
        })

    # -------------------------
    # Generate Excel
    # -------------------------
    df = pd.DataFrame(all_issues)

    os.makedirs(output_folder, exist_ok=True)
    output_path = os.path.join(output_folder, "G2_5_Performance_Constraints.xlsx")
    df.to_excel(output_path, index=False)

    print("G2.5 check completed")
    print("Excel generated at:", output_path)

    return df, output_path

def G2_6_logic():
    srs_path = str( G2Config.G2_SRS)
    g2_6_output_dir = str(G2Config.OUTPUT_DIR)

    excel_path, df = check_g2_6_and_generate_excel(srs_path, g2_6_output_dir)

    if df.empty:
        print("G2.6 PASS – No ambiguous timing terms found.")
    else:
        print("G2.6 FAIL – Findings detected.")
        print("Report generated at:", excel_path)

def run_g2_4_7():
    # ===== PATHS =====
    # Get project root dynamically (Review_Framework folder)
    PROJECT_ROOT = Path(__file__).resolve().parent.parent

    INPUT_DIR = PROJECT_ROOT / "inputs"
    OUTPUT_DIR = PROJECT_ROOT / "outputs"
    OUTPUT_FILE = OUTPUT_DIR / "CI_G2_4_7_output.xlsx"

    # Create outputs folder if not exists
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # ================= READ INPUTS =================
    # Expecting collect_requirements(...) -> List[Tuple[req_id, text]] or similar
    collected = io_utils.collect_requirements(str(INPUT_DIR))

    # Normalize to a list of (rid, text). If collect_requirements returned raw text, use extractor.
    if isinstance(collected, str):
        # collected is a single blob of text
        requirements = g2.g2_4_7_logic.extract_requirements(collected)
    else:
        # Assume already in [(rid, text), ...] form; if it's dict, convert to tuples
        if isinstance(collected, dict):
            requirements = list(collected.items())
        else:
            requirements = collected

    # ===== CREATE EXCEL =====
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirement Quality"

    # Title
    ws.merge_cells("A1:F1")
    ws["A1"] = "High-Level requirements are accurate and consistent"
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    # ===== UPDATED HEADER ORDER =====
    headers = [
        "Requirement ID",
        "G 2.4.1:Derived Requirement",
        "G 2.4.2:Derived Req Justified",
        "G 2.4.3:Justification Appropriate",
        "G 2.7.1:Sufficiently Detailed",
        "G 2.7.2:Review Reason"
    ]

    ws.append(headers)

    for col in range(1, len(headers) + 1):
        ws.cell(row=2, column=col).font = Font(bold=True)

    # ===== Populate rows with UPDATED order =====
    for rid, text in requirements:
        detail_status, detail_reason = g2.g2_4_7_logic.check_sufficient_detail(text)
        derived_status, justified_status, justification_quality = \
            g2.g2_4_7_logic.check_derived_requirement(text)

        ws.append([
            rid,
            derived_status,
            justified_status,
            justification_quality,
            detail_status,
            detail_reason
        ])

    # ===== Filter =====
    ws.auto_filter.ref = f"A2:F{ws.max_row}"

    # ===== Auto width safe =====
    for col in ws.columns:
        max_length = 0
        column_letter = None

        for cell in col:
            if isinstance(cell, MergedCell):
                continue

            if column_letter is None:
                column_letter = cell.column_letter

            if cell.value:
                max_length = max(max_length, len(str(cell.value)))

        if column_letter:
            ws.column_dimensions[column_letter].width = max_length + 3

    wb.save(OUTPUT_FILE)

    print("G2.4.7: ✅ Derived + Justification Quality verification completed")
    print(f"📄 Output saved at: {OUTPUT_FILE}")


def G2_logic():

    # To check G2.2 Sub Point:
    G2_2_logic()

    # To check G2.3 Sub Point:
    G2_3_logic()

    #To check G2.5 Sub Point:
    G2_5_logic(str(G2Config.G2_SRS) , str(CommonConfig.HARDWARE_DS_FILE),str(G2Config.OUTPUT_DIR))

    # To check G2.6 Sub Point:
    G2_6_logic()

    run_g2_4_7()

    combine_g2_excels(str(G2Config.OUTPUT_DIR))


