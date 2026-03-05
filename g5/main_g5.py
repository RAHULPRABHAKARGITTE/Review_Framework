from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

# from config import (
#     INPUT_SRS_FILE,
#     OUTPUT_REPORT_FILE,
#     REPORT_TITLE,
#     COLUMN_HEADERS,
#     STANDARDS_FILE,
#     AMBIGUOUS_WORDS,
#     FORBIDDEN_MODALS,
#     SAFETY_KEYWORDS,
#     MITIGATION_KEYWORDS,
#     FORBIDDEN_VAGUE,
#     PASSIVE_PATTERNS,
#     SUBJECT_PREFIX,
#     MAX_SHALL_COUNT,
#     STOP_WORDS
# )
from config import G5Config

from io_utils import _extract_requirements_from_docx
from g5.g5_logic import run_g5_checks


# ============================================================
# Keyword + Guideline Loader
# ============================================================
def load_standards():
    base_dir = Path(__file__).resolve().parent.parent
    docx_path = base_dir / "inputs" / G5Config.STANDARDS_FILE

    data = {
        "keywords": {
            "ambiguous_words": [],
            "forbidden_modals": [],
            "safety_keywords": [],
            "mitigation_keywords": [],
            "forbidden_vague": [],
            "passive_patterns": [],
            "stop_words": [],
            "subject_prefix": G5Config.SUBJECT_PREFIX,
            "max_shall_count": G5Config.MAX_SHALL_COUNT,
        },
        "parameters": {
            "expected_id_prefix": "SCU_STC_SRS_",
            "input_srs_file": G5Config.INPUT_SRS_FILE,
            "standards_file": G5Config.STANDARDS_FILE,
            "output_report_file": G5Config.OUTPUT_REPORT_FILE,
        },
        "guidelines": {}
    }

    if docx_path.exists():
        print(f"✅ Using standards from: {docx_path}")
        doc = Document(docx_path)

        headings = [para.text.strip().lower() for para in doc.paragraphs if para.text.strip().lower().startswith("table")]
        heading_iter = iter(headings)

        for table in doc.tables:
            heading_text = next(heading_iter, "").lower()
            category = None
            if "safety keyword" in heading_text:
                category = "safety_keywords"
            elif "mitigation keyword" in heading_text:
                category = "mitigation_keywords"
            elif "weak modal" in heading_text:
                category = "forbidden_modals"
            elif "vague" in heading_text:
                category = "forbidden_vague"
            elif "stop word" in heading_text:
                category = "stop_words"
            elif "passive pattern" in heading_text:
                category = "passive_patterns"
            elif "ambiguous word" in heading_text:
                category = "ambiguous_words"
            elif "format and structure" in heading_text:
                category = "parameters"

            values = []
            for row in table.rows:
                texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if category == "parameters" and len(texts) == 2:
                    key = texts[0].lower().replace(" ", "_")
                    value = texts[1]
                    if key in data["parameters"]:
                        data["parameters"][key] = value
                else:
                    values.extend(texts)

            if category and category in data["keywords"]:
                data["keywords"][category].extend(values)

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if "Safety Requirement" in text:
                data["guidelines"]["5.1"] = text
            elif "Single Functionality" in text:
                data["guidelines"]["5.2"] = text
            elif "Project Requirement" in text:
                data["guidelines"]["5.3"] = text
            elif "Ambiguous Word" in text:
                data["guidelines"]["5.4"] = text
            elif "Format and Structural" in text:
                data["guidelines"]["5.5"] = text

    else:
        print(f"⚠ File not found at {docx_path}. Using config.py defaults.")
        data["keywords"].update({
            "ambiguous_words": G5Config.AMBIGUOUS_WORDS,
            "forbidden_modals": G5Config.FORBIDDEN_MODALS,
            "forbidden_vague": G5Config.FORBIDDEN_VAGUE,
            "safety_keywords": G5Config.SAFETY_KEYWORDS,
            "mitigation_keywords": G5Config.MITIGATION_KEYWORDS,
            "passive_patterns": G5Config.PASSIVE_PATTERNS,
            "stop_words": G5Config.STOP_WORDS,
            "subject_prefix": G5Config.SUBJECT_PREFIX,
            "max_shall_count": G5Config.MAX_SHALL_COUNT,
        })
        data["parameters"].update({
            "expected_id_prefix": "SCU_STC_SRS_",
            "input_srs_file": G5Config.INPUT_SRS_FILE,
            "standards_file": G5Config.STANDARDS_FILE,
            "output_report_file": G5Config.OUTPUT_REPORT_FILE,
        })
        data["guidelines"].update({
            "5.1": "Safety requirements must include mitigation and avoid vague wording",
            "5.2": "Requirement shall contain only one 'shall'",
            "5.3": "Requirement must start with subject prefix",
            "5.4": "Requirement shall not contain ambiguous wording",
            "5.5": "Requirement IDs must follow project format",
        })

    return data


# ============================================================
# Excel Report Generator
# ============================================================
def generate_excel_report(all_req_ids, findings_by_req, output_path, guidelines, keywords, params):
    wb = Workbook()
    ws = wb.active
    ws.title = "SRS Review"

    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(G5Config.COLUMN_HEADERS))
    ws["A1"] = G5Config.REPORT_TITLE
    ws["A1"].font = Font(bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.append(G5Config.COLUMN_HEADERS)
    for col in range(1, len(G5Config.COLUMN_HEADERS) + 1):
        ws.cell(row=2, column=col).font = Font(bold=True)
        ws.cell(row=2, column=col).alignment = Alignment(wrap_text=True)

    row = 3
    for req_id in sorted(all_req_ids):
        issues = findings_by_req.get(req_id, [])

        status = { "5.1": "PASS", "5.2": "PASS", "5.3": "PASS", "5.4": "PASS", "5.5": "PASS" }
        failure_reasons = []

        for issue in issues:
            failure_reasons.append(issue)
            for g in status.keys():
                if g in issue:
                    status[g] = "FAIL"

        ws.append([
            req_id,
            status["5.1"], status["5.2"], status["5.3"], status["5.4"], status["5.5"],
            "\n".join(dict.fromkeys(failure_reasons)) if failure_reasons else "N/A"
        ])

        ws.cell(row=row, column=7).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    # Append summary section
    row += 2
    ws.cell(row=row, column=1, value="Summary of Extracted Standards").font = Font(bold=True)
    row += 1

    ws.cell(row=row, column=1, value="Keywords").font = Font(bold=True)
    row += 1
    for k, v in keywords.items():
        if isinstance(v, (list, set)):
            ws.cell(row=row, column=1, value=k)
            ws.cell(row=row, column=2, value=len(v))
        else:
            ws.cell(row=row, column=1, value=k)
            ws.cell(row=row, column=2, value=v)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Parameters").font = Font(bold=True)
    row += 1
    for k, v in params.items():
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        row += 1

    try:
        wb.save(output_path)
        print(f"\n✅ Excel report generated: {output_path}")
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback = output_path.with_stem(output_path.stem + "_" + timestamp)
        wb.save(fallback)
        print(f"\n⚠ File open. Saved as:\n{fallback}")

# ============================================================
# Main Orchestration
# ============================================================
def run_g5():
    standards = load_standards()
    keywords = standards["keywords"]
    params = standards["parameters"]
    guidelines = standards["guidelines"]

    base_dir = Path(__file__).resolve().parent.parent
    srs_path = base_dir / "inputs" / params["input_srs_file"]
    output_dir = base_dir / "outputs"
    output_dir.mkdir(exist_ok=True)
    output_path = output_dir / params["output_report_file"]

    # Extract requirements from the SRS file
    requirements = _extract_requirements_from_docx(srs_path)

    # Run guideline checks
    g5_results, all_findings = run_g5_checks(requirements, keywords)

    # Collect all requirement IDs
    valid_ids = {req["id"] for req in requirements if req.get("id")}
    reported_ids = {req_id for req_id, _ in all_findings if req_id}
    all_req_ids = valid_ids | reported_ids

    # Organize findings by requirement
    findings_by_req = defaultdict(list)
    for req_id, issue in all_findings:
        findings_by_req[req_id].append(issue)

    # Generate Excel report with summary
    generate_excel_report(all_req_ids, findings_by_req, output_path, guidelines, keywords, params)
